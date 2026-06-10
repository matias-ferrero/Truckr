"""
Converts features.typ into a formatted .xlsx matrix.
Usage: uv run --with openpyxl docs/scripts/features_typ2xlsx.py
Output: docs/artifacts/features.xlsx
"""

import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── font sizes ───────────────────────────────────────────────────────────────
# Tweak these to adjust the whole document proportionally.

SZ_CORNER   = 13   # "Features / Personas" corner cell
SZ_HEADER   = 12   # feature column headers
SZ_PERSONA  = 12   # persona name cells + Average label
SZ_STARS    = 16   # star rating cells
SZ_AVG      = 14   # average value cells

# ── paths ─────────────────────────────────────────────────────────────────────

REPO_ROOT = Path(__file__).resolve().parents[2]
FEATURES_TYP = REPO_ROOT / "docs/artifacts/features.typ"
OUT = REPO_ROOT / "docs/artifacts/features.xlsx"

# ── parser ────────────────────────────────────────────────────────────────────

def _hex(rgb_call: str) -> str:
    """Extract the 6-char hex from 'rgb("#RRGGBB")'."""
    m = re.search(r'rgb\("#([0-9A-Fa-f]{6})"\)', rgb_call)
    return m.group(1).upper() if m else "FFFFFF"


def parse_personas_color(text: str) -> str:
    m = re.search(r'#let personas-color\s*=\s*rgb\("#([0-9A-Fa-f]{6})"\)', text)
    return m.group(1).upper() if m else "C9DAF8"


def parse_data_font_color(text: str) -> str:
    m = re.search(r'#let data-font-color\s*=\s*rgb\("#([0-9A-Fa-f]{6})"\)', text)
    return m.group(1).upper() if m else "FBBC04"


def parse_avg_colors(text: str) -> list[str]:
    start = text.find("#let avg-colors = (")
    if start == -1:
        return []
    open_pos = text.index("(", start)
    depth, i = 0, open_pos
    while i < len(text):
        if text[i] == "(":
            depth += 1
        elif text[i] == ")":
            depth -= 1
            if depth == 0:
                break
        i += 1
    block = text[open_pos + 1 : i]
    return [_hex(chunk) for chunk in re.findall(r'rgb\("#[0-9A-Fa-f]{6}"\)', block)]


def parse_feature_names(text: str) -> list[str]:
    """
    Extract feature header labels from table.cell(align: center + horizon) blocks.
    Skips the first cell which is the "Features / Personas" corner cell.
    """
    # Each feature cell: table.cell(align: center + horizon)[\n  #set par...\n  *Name*\n]
    pattern = re.compile(
        r"table\.cell\(align:\s*center \+ horizon\)\[\s*#set par\(justify: false\)\s*\*([^*]+)\*",
        re.DOTALL,
    )
    raw = [m.group(1).strip() for m in pattern.finditer(text)]
    # Clean Typst escapes: \* → *, \ → space
    cleaned = []
    for name in raw:
        name = name.replace("\\ ", " ").replace("\\*", "*").replace("\\", "")
        cleaned.append(" ".join(name.split()))
    return cleaned


def parse_persona_names(text: str) -> list[str]:
    """
    Extract persona name labels from single-line table.cell(fill: personas-color)[Name] calls.
    Excludes the "Average" row and the multi-line corner header.
    """
    pattern = re.compile(r"table\.cell\(fill: personas-color\)\[([^\]]+)\]")
    names = [m.group(1).strip() for m in pattern.finditer(text)]
    return [n for n in names if n.lower() != "average"]


def parse_ratings(text: str) -> list[list[int]]:
    """Extract the ratings matrix as a list of rows."""
    start = text.find("#let ratings = (")
    if start == -1:
        return []
    open_pos = text.index("(", start)
    depth, i = 0, open_pos
    while i < len(text):
        if text[i] == "(":
            depth += 1
        elif text[i] == ")":
            depth -= 1
            if depth == 0:
                break
        i += 1
    block = text[open_pos + 1 : i]
    rows = []
    for row_m in re.finditer(r"\(([^)]+)\)", block):
        nums = [int(x.strip()) for x in row_m.group(1).split(",") if x.strip().isdigit()]
        if nums:
            rows.append(nums)
    return rows


def col_avg(ratings: list[list[int]], col: int) -> float:
    vals = [row[col] for row in ratings if col < len(row)]
    return sum(vals) / len(vals) if vals else 0.0


def avg_color_idx(avg: float, n_colors: int) -> int:
    """Same formula as the Typst source: min(n-1, floor((avg-1) * 3.5))."""
    return min(n_colors - 1, math.floor((avg - 1) * 3.5))


def stars(n: int) -> str:
    return "★" * n + "☆" * (5 - n)


# ── style helpers ─────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _border() -> Border:
    side = Side(style="thin", color="AAAAAA")
    return Border(left=side, right=side, top=side, bottom=side)


def _font(bold=False, color="000000", size=11, name="Calibri") -> Font:
    return Font(bold=bold, color=color, size=size, name=name)


def _align(h="center", v="center", wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


# ── workbook builder ──────────────────────────────────────────────────────────

def build_xlsx(
    feature_names: list[str],
    persona_names: list[str],
    ratings: list[list[int]],
    personas_color: str,
    data_font_color: str,
    avg_colors: list[str],
) -> None:
    n_features = len(feature_names)
    n_personas = len(persona_names)

    wb = Workbook()
    ws = wb.active
    ws.title = "Features Matrix"

    # ── row 1: header ─────────────────────────────────────────────────────────
    corner = ws.cell(row=1, column=1, value="Features / Personas")
    corner.fill = _fill(personas_color)
    corner.font = _font(bold=True, color="000000", size=SZ_CORNER)
    corner.alignment = _align(h="center", wrap=True)
    corner.border = _border()

    for col_idx, name in enumerate(feature_names, start=2):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = _fill("FFFFFF")
        cell.font = _font(bold=True, color="000000", size=SZ_HEADER)
        cell.alignment = _align(h="center", wrap=True)
        cell.border = _border()

    # ── rows 2…n+1: persona data ──────────────────────────────────────────────
    for row_idx, (persona, row_ratings) in enumerate(
        zip(persona_names, ratings), start=2
    ):
        # persona name cell
        name_cell = ws.cell(row=row_idx, column=1, value=persona)
        name_cell.fill = _fill(personas_color)
        name_cell.font = _font(bold=True, color="000000", size=SZ_PERSONA)
        name_cell.alignment = _align(h="left", v="center")
        name_cell.border = _border()

        for col_idx, rating in enumerate(row_ratings, start=2):
            cell = ws.cell(row=row_idx, column=col_idx, value=stars(rating))
            cell.fill = _fill("FFFFFF")
            cell.font = Font(
                bold=False,
                color=data_font_color,
                size=SZ_STARS,
                name="DejaVu Sans",
            )
            cell.alignment = _align(h="center")
            cell.border = _border()

    # ── last row: averages ────────────────────────────────────────────────────
    avg_row = n_personas + 2
    avg_label = ws.cell(row=avg_row, column=1, value="Average")
    avg_label.fill = _fill(personas_color)
    avg_label.font = _font(bold=True, color="000000", size=SZ_PERSONA)
    avg_label.alignment = _align(h="left", v="center")
    avg_label.border = _border()

    for col_idx in range(n_features):
        avg = col_avg(ratings, col_idx)
        idx = avg_color_idx(avg, len(avg_colors))
        bg = avg_colors[idx] if avg_colors else "F1C232"
        avg_val = round(avg, 1)
        display = f"{avg_val:.1f}"

        cell = ws.cell(row=avg_row, column=col_idx + 2, value=float(display))
        cell.fill = _fill(bg)
        cell.font = _font(bold=True, color="000000", size=SZ_AVG)
        cell.alignment = _align(h="center")
        cell.number_format = "0.0"
        cell.border = _border()

    # ── column widths and row heights ─────────────────────────────────────────
    ws.column_dimensions[get_column_letter(1)].width = 30
    for col_idx in range(2, n_features + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    ws.row_dimensions[1].height = 60  # header needs room for wrapped text
    for row_idx in range(2, avg_row + 1):
        ws.row_dimensions[row_idx].height = 28

    # ── freeze the header row and name column ─────────────────────────────────
    ws.freeze_panes = "B2"

    wb.save(OUT)
    print(f"Saved → {OUT}")


# ── entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    text = FEATURES_TYP.read_text(encoding="utf-8")

    feature_names = parse_feature_names(text)
    persona_names = parse_persona_names(text)
    ratings = parse_ratings(text)
    personas_color = parse_personas_color(text)
    data_font_color = parse_data_font_color(text)
    avg_colors = parse_avg_colors(text)

    assert len(feature_names) == 13, f"Expected 13 features, got {len(feature_names)}"
    assert len(persona_names) == len(ratings), (
        f"Persona count ({len(persona_names)}) != ratings rows ({len(ratings)})"
    )

    build_xlsx(feature_names, persona_names, ratings, personas_color, data_font_color, avg_colors)
