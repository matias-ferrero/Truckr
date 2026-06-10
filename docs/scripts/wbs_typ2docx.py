"""
Converts wbs.typ into a formatted hierarchical .docx outline.
Usage: uv run --with python-docx docs/scripts/wbs_typ2docx.py
Output: docs/artifacts/wbs.docx
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── paths ─────────────────────────────────────────────────────────────────────

REPO_ROOT = Path(__file__).resolve().parents[2]
WBS_TYP = REPO_ROOT / "docs/artifacts/wbs.typ"
OUT = REPO_ROOT / "docs/artifacts/wbs.docx"

# ── font sizes ───────────────────────────────────────────────────────────────
# Tweak these to adjust the whole document proportionally.

SZ_TITLE  = Pt(24)   # "WBS — Work Breakdown Structure" h1
SZ_ROOT   = Pt(16)   # root node banner ("Truckr® — Plataforma de Transporte")
SZ_L1     = Pt(15)   # level-1 section heading  (e.g. "1. Autenticación y Cuentas")
SZ_L2     = Pt(12)   # level-2 sub-section      (e.g. "1.1 Registro")
SZ_L3     = Pt(11)   # level-3 leaf item        (e.g. "1.1.1 Registrarse")
SZ_FOOTER = Pt(10)   # post-MVP footnote

# ── colours (brand palette) ───────────────────────────────────────────────────

C_BRAND       = RGBColor(0x15, 0x43, 0x60)  # #154360 — deep navy
C_BRAND_MID   = RGBColor(0x1F, 0x61, 0x8D)  # #1F618D — mid blue
C_BRAND_LIGHT = RGBColor(0xD6, 0xEA, 0xF8)  # #D6EAF8 — pale blue (shading)
C_POST_MVP    = RGBColor(0x88, 0x88, 0x88)  # grey for post-MVP labels

# ── indentation per level ─────────────────────────────────────────────────────

INDENT = {1: Cm(0.0), 2: Cm(0.8), 3: Cm(1.6)}


# ── parser ────────────────────────────────────────────────────────────────────

# Matches: node(<var>, <hw>, <hh>, [<label>])
# Captures the label text inside the brackets.
_NODE_RE = re.compile(
    r"node\(\s*[\w-]+\s*,\s*[\w]+\s*,\s*[\w]+\s*,\s*\[([^\]]*)\]",
)


def _clean(raw: str) -> str:
    """Strip Typst markup; return plain text suitable for a Word document."""
    label = raw.replace("\\ ", " ")       # soft line-break → space
    label = label.replace("\\*", "*")     # escaped asterisk → literal *
    label = label.replace("\\", "")       # drop remaining escapes
    label = re.sub(r"\*([^*]+)\*", r"\1", label)  # *bold* → plain
    return " ".join(label.split()).strip()


def _is_post_mvp(raw: str) -> bool:
    """True when the Typst label carries a post-MVP \\* marker."""
    return "\\*" in raw


def _level(label: str) -> int:
    """Determine WBS depth from the numeric prefix of the cleaned label."""
    if re.match(r"^\d+\.\d+\.\d+", label):
        return 3
    if re.match(r"^\d+\.\d+", label):
        return 2
    if re.match(r"^\d+\.", label):
        return 1
    return 0  # root node


def parse_wbs(path: Path) -> list[dict]:
    """Return nodes in file order (which is DFS pre-order in the WBS)."""
    nodes = []
    for m in _NODE_RE.finditer(path.read_text(encoding="utf-8")):
        raw = m.group(1)
        label = _clean(raw)
        nodes.append({"label": label, "post_mvp": _is_post_mvp(raw)})
    return nodes


# ── docx helpers ──────────────────────────────────────────────────────────────

def _shade(paragraph, fill_hex: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def _post_mvp_tag(paragraph) -> None:
    r = paragraph.add_run("  (Post-MVP)")
    r.font.size = Pt(9)
    r.font.color.rgb = C_POST_MVP
    r.italic = True


def _add_node(doc: Document, label: str, post_mvp: bool, level: int) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format

    if level == 0:
        _shade(p, "154360")
        pf.left_indent = Cm(0.3)
        pf.space_before = Pt(4)
        pf.space_after = Pt(6)
        run = p.add_run(label)
        run.bold = True
        run.font.size = SZ_ROOT
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    elif level == 1:
        _shade(p, "D6EAF8")
        pf.left_indent = INDENT[1]
        pf.space_before = Pt(10)
        pf.space_after = Pt(2)
        run = p.add_run(label)
        run.bold = True
        run.font.size = SZ_L1
        run.font.color.rgb = C_BRAND
        if post_mvp:
            _post_mvp_tag(p)

    elif level == 2:
        pf.left_indent = INDENT[2]
        pf.space_before = Pt(5)
        pf.space_after = Pt(2)
        run = p.add_run(label)
        run.bold = True
        run.font.size = SZ_L2
        run.font.color.rgb = C_BRAND_MID
        if post_mvp:
            _post_mvp_tag(p)

    else:  # level 3
        pf.left_indent = INDENT[3]
        pf.first_line_indent = Cm(-0.4)
        pf.space_before = Pt(2)
        pf.space_after = Pt(1)
        bullet = p.add_run("• ")
        bullet.font.size = SZ_L3
        bullet.font.color.rgb = C_BRAND_MID
        run = p.add_run(label)
        run.font.size = SZ_L3
        if post_mvp:
            run.italic = True
            run.font.color.rgb = C_POST_MVP
        else:
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


# ── document builder ──────────────────────────────────────────────────────────

def build_docx(nodes: list[dict]) -> None:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = doc.add_heading("WBS — Work Breakdown Structure", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        run.font.size = SZ_TITLE
        run.font.color.rgb = C_BRAND

    doc.add_paragraph()

    for node in nodes:
        _add_node(doc, node["label"], node["post_mvp"], _level(node["label"]))

    doc.add_paragraph()
    footer = doc.add_paragraph()
    r = footer.add_run(
        "(*) Estos trabajos han sido tomados para el momento post-MVP, "
        "por lo que no se verán implementados en el producto mínimo viable."
    )
    r.font.size = SZ_FOOTER
    r.italic = True
    r.font.color.rgb = C_POST_MVP

    doc.save(OUT)
    print(f"Saved → {OUT}")


if __name__ == "__main__":
    build_docx(parse_wbs(WBS_TYP))
