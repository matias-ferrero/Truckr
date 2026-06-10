"""
Converts backlog-us .typ source files into a formatted .docx.
Usage: uv run --with python-docx docs/scripts/backlog_typ2docx.py
Output: docs/artifacts/backlog-us.docx
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── font sizes ───────────────────────────────────────────────────────────────
# Tweak these to adjust the whole document proportionally.

SZ_DOC_TITLE    = Pt(24)   # "Backlog — User Stories" h1
SZ_SUBTITLE     = Pt(11)   # intro italic line
SZ_BANNER       = Pt(16)   # release banner (MVP / Post-MVP)
SZ_US_TITLE     = Pt(15)   # each US heading (h2)
SZ_SECTION_LABEL = Pt(12)  # "Descripción" / "Criterios de Aceptación" labels
SZ_BODY         = Pt(11)   # description text, AC items
SZ_META         = Pt(10)   # Release / Prioridad / Épica line (secondary info)

# ── paths ────────────────────────────────────────────────────────────────────

REPO_ROOT = Path(__file__).resolve().parents[2]
US_DIR = REPO_ROOT / "docs/artifacts/backlog-us"
BACKLOG_TYP = REPO_ROOT / "docs/artifacts/backlog-us.typ"
OUT = REPO_ROOT / "docs/artifacts/backlog-us.docx"

# ── release order from backlog-us.typ ────────────────────────────────────────

RELEASES = [
    {
        "label": "MVP — Release 1",
        "color": RGBColor(0x6A, 0xA8, 0x4F),
        "files": [],
    },
    {
        "label": "Post MVP — Release 2",
        "color": RGBColor(0x3D, 0x7A, 0xB5),
        "files": [],
    },
    {
        "label": "Post MVP — Release 3",
        "color": RGBColor(0x7B, 0x5E, 0xA7),
        "files": [],
    },
]


def load_release_order():
    """Read backlog-us.typ and collect the include order per release."""
    text = BACKLOG_TYP.read_text(encoding="utf-8")
    current = None
    for line in text.splitlines():
        if "MVP — Release 1" in line and "release-banner" in line:
            current = 0
        elif "Release 2" in line and "release-banner" in line:
            current = 1
        elif "Release 3" in line and "release-banner" in line:
            current = 2
        elif line.strip().startswith('#us("backlog-us/'):
            m = re.search(r'"backlog-us/(US\d+\.typ)"', line)
            if m and current is not None:
                RELEASES[current]["files"].append(US_DIR / m.group(1))


# ── parser ────────────────────────────────────────────────────────────────────

def parse_us(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")

    def field(key):
        m = re.search(rf"\*{key}:\*\s*(.+?)(?:\s*\\\\|$)", text, re.MULTILINE)
        return m.group(1).strip() if m else ""

    # title: == US7: Ofertar Retiro …
    title_m = re.match(r"==\s+(.+)", text.strip())
    title = title_m.group(1).strip() if title_m else path.stem

    # description block
    desc_m = re.search(r"\*Descripción:\*\s*\n(.*?)\n\n", text, re.DOTALL)
    description = desc_m.group(1).strip() if desc_m else ""

    # acceptance criteria: numbered list items starting with +
    ac_block_m = re.search(r"\*Criterios de Aceptación:\*\s*\n(.*?)(?:\n\n|$)", text, re.DOTALL)
    criteria = []
    if ac_block_m:
        for item in re.findall(r"^\s*\+\s*(.+)", ac_block_m.group(1), re.MULTILINE):
            criteria.append(item.strip())
        # sub-items (indented with -)
        # rebuild with sub-items attached to previous criterion
        criteria = []
        current_item = None
        for line in ac_block_m.group(1).splitlines():
            top = re.match(r"^\s*\+\s*(.+)", line)
            sub = re.match(r"^\s+-\s*(.+)", line)
            if top:
                if current_item is not None:
                    criteria.append(current_item)
                current_item = {"text": top.group(1).strip(), "subs": []}
            elif sub and current_item is not None:
                current_item["subs"].append(sub.group(1).strip())
        if current_item is not None:
            criteria.append(current_item)

    return {
        "title": title,
        "release": field("Release"),
        "priority": field("Prioridad"),
        "epic": field("Épica"),
        "description": description,
        "criteria": criteria,
    }


# ── docx helpers ──────────────────────────────────────────────────────────────

def set_paragraph_shading(paragraph, fill_hex: str):
    """Apply background color to a paragraph via XML."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def add_release_banner(doc: Document, label: str, color: RGBColor):
    doc.add_paragraph()  # breathing room before banner
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hex_color = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
    set_paragraph_shading(p, hex_color)
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.left_indent = Cm(0.3)
    run = p.add_run(label)
    run.bold = True
    run.font.size = SZ_BANNER
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_us(doc: Document, us: dict):
    paragraphs = []

    # US heading
    heading = doc.add_heading(us["title"], level=2)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(4)
    for run in heading.runs:
        run.font.size = SZ_US_TITLE
    paragraphs.append(heading)

    # metadata row
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(3)
    for label, value in [
        ("Release", us["release"]),
        ("Prioridad", us["priority"]),
        ("Épica", us["epic"]),
    ]:
        run_label = meta.add_run(f"{label}: ")
        run_label.bold = True
        run_label.font.size = SZ_META
        run_label.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        run_value = meta.add_run(value + "   ")
        run_value.font.size = SZ_META
        run_value.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    paragraphs.append(meta)

    # description
    if us["description"]:
        label_p = doc.add_paragraph()
        label_p.paragraph_format.space_before = Pt(8)
        label_p.paragraph_format.space_after = Pt(3)
        r = label_p.add_run("Descripción")
        r.bold = True
        r.font.size = SZ_SECTION_LABEL
        paragraphs.append(label_p)

        desc_p = doc.add_paragraph(us["description"])
        desc_p.paragraph_format.left_indent = Cm(0.5)
        desc_p.paragraph_format.space_after = Pt(6)
        desc_p.runs[0].font.size = SZ_BODY
        desc_p.runs[0].font.italic = True
        paragraphs.append(desc_p)

    # acceptance criteria
    if us["criteria"]:
        label_p = doc.add_paragraph()
        label_p.paragraph_format.space_before = Pt(8)
        label_p.paragraph_format.space_after = Pt(3)
        r = label_p.add_run("Criterios de Aceptación")
        r.bold = True
        r.font.size = SZ_SECTION_LABEL
        paragraphs.append(label_p)

        for i, item in enumerate(us["criteria"], 1):
            text = item["text"] if isinstance(item, dict) else item
            subs = item["subs"] if isinstance(item, dict) else []

            # manual numbering so each US resets to 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(3)
            num_run = p.add_run(f"{i}. ")
            num_run.bold = True
            num_run.font.size = SZ_BODY
            text_run = p.add_run(text)
            text_run.font.size = SZ_BODY
            paragraphs.append(p)

            for sub in subs:
                sp = doc.add_paragraph()
                sp.paragraph_format.left_indent = Cm(1.4)
                sp.paragraph_format.first_line_indent = Cm(-0.4)
                sp.paragraph_format.space_after = Pt(3)
                bullet_run = sp.add_run("• ")
                bullet_run.font.size = SZ_BODY
                sr = sp.add_run(sub)
                sr.font.size = SZ_BODY
                paragraphs.append(sp)

    # separator line
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(8)
    sep.paragraph_format.space_after = Pt(0)
    pPr = sep._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)
    paragraphs.append(sep)

    # keep all paragraphs in this US together: if the block doesn't fit on the
    # current page, Word moves everything to the next page instead of splitting
    for p in paragraphs[:-1]:  # all except the separator
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True


# ── document setup ────────────────────────────────────────────────────────────

def build_docx():
    load_release_order()

    doc = Document()

    # page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # document title
    title = doc.add_heading("Backlog — User Stories", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in title.runs:
        run.font.size = SZ_DOC_TITLE

    subtitle = doc.add_paragraph(
        "Todas las historias siguen el framework de las 3 C's (Card, Conversation, Confirmation) "
        "y los criterios INVEST. Prioridades: Alta / Media / Baja."
    )
    subtitle.runs[0].font.size = SZ_SUBTITLE
    subtitle.runs[0].italic = True
    subtitle.paragraph_format.space_after = Pt(6)

    for release in RELEASES:
        if not release["files"]:
            continue
        add_release_banner(doc, release["label"], release["color"])
        for f in release["files"]:
            if f.exists():
                add_us(doc, parse_us(f))
            else:
                print(f"  WARNING: {f.name} not found, skipping")

    doc.save(OUT)
    print(f"Saved → {OUT}")


if __name__ == "__main__":
    build_docx()
